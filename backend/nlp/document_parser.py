import hashlib
import io
from pathlib import Path
from typing import Tuple
import pymupdf as fitz  # PyMuPDF
import pdfplumber
from docx import Document
from backend.config import settings


class ResumeParseError(Exception):
    def __init__(self, message: str, error_code: str = "PARSE_ERROR"):
        super().__init__(message)
        self.message = message
        self.error_code = error_code


def compute_sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def validate_file_metadata(filename: str, content: bytes) -> str:
    if not filename or "." not in filename:
        raise ResumeParseError(
            f"Invalid file '{filename}': Missing file extension. Supported formats: PDF, DOCX, TXT.",
            error_code="INVALID_FILE",
        )
    ext = filename.rsplit(".", 1)[-1].lower().strip()
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise ResumeParseError(
            f"Unsupported file format '.{ext}' for '{filename}'. Please upload PDF, DOCX, or TXT files only.",
            error_code="UNSUPPORTED_FORMAT",
        )
    if len(content) == 0:
        raise ResumeParseError(
            f"The uploaded file '{filename}' is empty (0 bytes).",
            error_code="EMPTY_FILE",
        )
    if len(content) > settings.MAX_UPLOAD_SIZE_BYTES:
        size_mb = round(len(content) / (1024 * 1024), 2)
        raise ResumeParseError(
            f"File '{filename}' ({size_mb} MB) exceeds the maximum allowed size of {settings.MAX_UPLOAD_SIZE_MB} MB.",
            error_code="FILE_TOO_LARGE",
        )
    return ext


def extract_text_from_pdf(content: bytes, filename: str) -> str:
    # Primary extraction via PyMuPDF (fitz)
    text_parts = []
    fitz_error = None
    try:
        with fitz.open(stream=content, filetype="pdf") as doc:
            if doc.is_encrypted:
                raise ResumeParseError(
                    f"PDF file '{filename}' is password-protected/encrypted and cannot be parsed.",
                    error_code="CORRUPTED_PDF",
                )
            for page in doc:
                page_text = page.get_text("text")
                if page_text:
                    text_parts.append(page_text)
    except ResumeParseError:
        raise
    except Exception as e:
        fitz_error = e

    extracted = "\n".join(text_parts).strip()
    if len(extracted) >= 25:
        return extracted

    # Fallback extraction via pdfplumber
    try:
        plumber_parts = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                pt = page.extract_text()
                if pt:
                    plumber_parts.append(pt)
        extracted_plumber = "\n".join(plumber_parts).strip()
        if extracted_plumber:
            return extracted_plumber
    except Exception as plumber_err:
        if fitz_error:
            raise ResumeParseError(
                f"Corrupted or unreadable PDF file '{filename}'. Details: {str(fitz_error)}",
                error_code="CORRUPTED_PDF",
            ) from plumber_err

    if fitz_error:
        raise ResumeParseError(
            f"Corrupted or invalid PDF file '{filename}'. Unable to read PDF stream.",
            error_code="CORRUPTED_PDF",
        )
    return extracted


def extract_text_from_docx(content: bytes, filename: str) -> str:
    try:
        doc = Document(io.BytesIO(content))
        parts = []
        for para in doc.paragraphs:
            if para.text and para.text.strip():
                parts.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if row_cells:
                    parts.append(" | ".join(row_cells))
        return "\n".join(parts).strip()
    except Exception as e:
        raise ResumeParseError(
            f"Corrupted or unreadable DOCX file '{filename}'. Details: {str(e)}",
            error_code="CORRUPTED_DOCX",
        ) from e


def extract_text_from_txt(content: bytes, filename: str) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return content.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    raise ResumeParseError(
        f"Unable to decode text file '{filename}'. Ensure it is valid UTF-8 or ASCII text.",
        error_code="INVALID_TEXT_ENCODING",
    )


def parse_resume_document(filename: str, content: bytes) -> Tuple[str, str, str]:
    """
    Validates and extracts text from PDF, DOCX, or TXT bytes.
    Returns (extracted_text, file_extension, sha256_hash).
    """
    ext = validate_file_metadata(filename, content)

    if ext == "pdf":
        raw_text = extract_text_from_pdf(content, filename)
    elif ext == "docx":
        raw_text = extract_text_from_docx(content, filename)
    elif ext == "txt":
        raw_text = extract_text_from_txt(content, filename)
    else:
        raise ResumeParseError(
            f"Unsupported file extension '.{ext}'.",
            error_code="UNSUPPORTED_FORMAT",
        )

    cleaned = raw_text.strip()
    if len(cleaned) < 20:
        raise ResumeParseError(
            f"Resume '{filename}' contains no readable text or is nearly empty. Please upload a valid text-based resume.",
            error_code="EMPTY_RESUME",
        )

    content_hash = compute_sha256(cleaned.encode("utf-8"))
    return cleaned, ext, content_hash
