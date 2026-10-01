import csv
import io
from datetime import datetime
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from backend.nlp.fairness_guard import FAIRNESS_DISCLAIMER


def generate_csv_export(job_info: Dict[str, Any], rows: List[Dict[str, Any]]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "Rank",
            "Candidate Name",
            "Email",
            "Location",
            "Overall Score (%)",
            "Recommendation",
            "Recruiter Decision",
            "Skills Score (%)",
            "Experience Score (%)",
            "Education Score (%)",
            "Experience (Years)",
            "Highest Education",
            "Candidate Skills",
            "Matched Requirements / Skills",
            "Missing Requirements / Skills",
            "Recruiter Notes",
        ]
    )

    for r in rows:
        writer.writerow(
            [
                f"#{r['rank']}",
                r["full_name"],
                r["email"],
                r["location"],
                f"{r['overall_score']:.1f}%",
                r["recommendation"],
                r["recruiter_decision"],
                f"{r['skills_percentage']:.1f}%",
                f"{r['experience_percentage']:.1f}%",
                f"{r['education_percentage']:.1f}%",
                r["total_experience_years"],
                r["highest_education"],
                ", ".join(r["all_skills"]),
                ", ".join(r["matched_skills"]),
                ", ".join(r["missing_skills"]),
                " | ".join(r["recruiter_notes"]) if r["recruiter_notes"] else "No notes",
            ]
        )

    return output.getvalue().encode("utf-8-sig")


def generate_excel_export(job_info: Dict[str, Any], rows: List[Dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidate Screening Results"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    sub_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    bold_font = Font(name="Calibri", size=10, bold=True, color="1E293B")
    regular_font = Font(name="Calibri", size=10, color="1E293B")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Report Header
    ws["A1"] = f"AI Resume Screening Report — {job_info['title']}"
    ws["A1"].font = title_font
    ws["A2"] = (
        f"Department: {job_info['department']} | Location: {job_info['location']} | "
        f"Min Experience: {job_info['min_experience_years']} yrs | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    ws["A2"].font = bold_font
    ws["A3"] = f"Fairness & Governance Notice: {FAIRNESS_DISCLAIMER}"
    ws["A3"].font = Font(name="Calibri", size=9, italic=True, color="475569")

    headers = [
        "Rank",
        "Candidate",
        "Overall Score",
        "Recommendation",
        "Recruiter Status",
        "Skills (%)",
        "Experience (%)",
        "Education (%)",
        "Years Exp",
        "Education",
        "All Extracted Skills",
        "Matched Requirements",
        "Missing Requirements",
        "Recruiter Notes",
    ]

    header_row = 5
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for idx, r in enumerate(rows, start=header_row + 1):
        vals = [
            f"#{r['rank']}",
            f"{r['full_name']} ({r['email']})",
            f"{r['overall_score']:.1f}%",
            r["recommendation"],
            r["recruiter_decision"],
            f"{r['skills_percentage']:.1f}%",
            f"{r['experience_percentage']:.1f}%",
            f"{r['education_percentage']:.1f}%",
            r["total_experience_years"],
            r["highest_education"],
            ", ".join(r["all_skills"]),
            ", ".join(r["matched_skills"]),
            ", ".join(r["missing_skills"]) or "None",
            " | ".join(r["recruiter_notes"]) if r["recruiter_notes"] else "—",
        ]
        for col_idx, val in enumerate(vals, start=1):
            c = ws.cell(row=idx, column=col_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if idx % 2 == 1:
                c.fill = sub_fill

    col_widths = [8, 28, 14, 18, 18, 13, 15, 14, 11, 16, 36, 34, 28, 36]
    for i, w in enumerate(col_widths, start=1):
        col_letter = ws.cell(row=header_row, column=i).column_letter
        ws.column_dimensions[col_letter].width = w

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def generate_pdf_export(job_info: Dict[str, Any], rows: List[Dict[str, Any]]) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=0.4 * inch,
        rightMargin=0.4 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.45 * inch,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0F172A"),
        spaceAfter=4,
    )
    sub_style = ParagraphStyle(
        "ReportSub",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#475569"),
        spaceAfter=6,
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1E293B"),
    )
    header_cell_style = ParagraphStyle(
        "TableHeaderCell",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.white,
        fontName="Helvetica-Bold",
    )

    elements = []
    elements.append(
        Paragraph(
            f"TalentPulse AI — Candidate Intelligence Report: {job_info['title']}",
            title_style,
        )
    )
    elements.append(
        Paragraph(
            f"<b>Department:</b> {job_info['department']} &nbsp;|&nbsp; "
            f"<b>Location:</b> {job_info['location']} &nbsp;|&nbsp; "
            f"<b>Min Experience:</b> {job_info['min_experience_years']} yrs &nbsp;|&nbsp; "
            f"<b>Total Candidates Ranked:</b> {len(rows)}",
            sub_style,
        )
    )
    elements.append(
        Paragraph(
            f"<i>Fairness & Ethical AI Disclaimer: {FAIRNESS_DISCLAIMER}</i>",
            sub_style,
        )
    )
    elements.append(Spacer(1, 8))

    table_data = [
        [
            Paragraph("Rank", header_cell_style),
            Paragraph("Candidate", header_cell_style),
            Paragraph("Match Score", header_cell_style),
            Paragraph("Breakdown (S/Ex/Ed)", header_cell_style),
            Paragraph("Exp & Edu", header_cell_style),
            Paragraph("Matched Skills", header_cell_style),
            Paragraph("Missing Skills", header_cell_style),
            Paragraph("Recruiter Notes & Status", header_cell_style),
        ]
    ]

    for r in rows:
        notes_str = "; ".join(r["recruiter_notes"]) if r["recruiter_notes"] else "No notes added"
        table_data.append(
            [
                Paragraph(f"<b>#{r['rank']}</b>", cell_style),
                Paragraph(
                    f"<b>{r['full_name']}</b><br/>{r['email']}<br/>{r['location']}",
                    cell_style,
                ),
                Paragraph(
                    f"<b>{r['overall_score']:.0f}%</b><br/>{r['recommendation']}",
                    cell_style,
                ),
                Paragraph(
                    f"Skills: {r['skills_percentage']:.0f}%<br/>"
                    f"Exp: {r['experience_percentage']:.0f}%<br/>"
                    f"Edu: {r['education_percentage']:.0f}%",
                    cell_style,
                ),
                Paragraph(
                    f"{r['total_experience_years']} yrs<br/>{r['highest_education']}",
                    cell_style,
                ),
                Paragraph(", ".join(r["matched_skills"][:10]) or "None", cell_style),
                Paragraph(", ".join(r["missing_skills"][:8]) or "None", cell_style),
                Paragraph(
                    f"<b>[{r['recruiter_decision']}]</b><br/>{notes_str[:160]}",
                    cell_style,
                ),
            ]
        )

    col_widths = [
        0.5 * inch,
        1.55 * inch,
        0.95 * inch,
        1.15 * inch,
        0.95 * inch,
        1.9 * inch,
        1.45 * inch,
        2.35 * inch,
    ]
    tbl = Table(table_data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F8FAFC")],
                ),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    elements.append(tbl)
    doc.build(elements)
    return buffer.getvalue()
